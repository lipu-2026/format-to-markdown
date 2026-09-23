import io
import sys
import threading
import unittest
import urllib.error
import urllib.request
import json
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app


def convert(data, name):
    return app.convert_file(data, name, include_metadata=False, embed_images=False)


def document_bytes(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class ConversionTests(unittest.TestCase):
    def test_corrupt_documents_are_never_successful_plain_text(self):
        for name in ['bad.pdf', 'bad.docx', 'bad.pptx', 'bad.xlsx', 'bad.epub', 'bad.zip', 'bad.png']:
            with self.subTest(name=name), self.assertRaises(ValueError):
                convert(b'this is not a document', name)

    def test_unknown_binary_is_rejected_but_utf8_is_retained(self):
        with self.assertRaises(ValueError):
            convert(b'\x00ABC\xff', 'data.bin')
        result = convert('中文内容'.encode(), 'notes.custom')
        self.assertIn('中文内容', result['markdown'])
        self.assertTrue(result['warning'])

    def test_markdown_hard_break_and_nested_fences(self):
        source = '# 标题\n\n第一行  \n第二行\n'
        self.assertEqual(convert(source.encode(), '测试.md')['markdown'], source)
        result = convert(b'print("```")', 'nested.py')['markdown']
        self.assertIn('````python', result)
        self.assertTrue(result.endswith('````\n'))

    def test_docx_keeps_block_order_and_table(self):
        from docx import Document
        document = Document()
        document.add_heading('资料标题', 1)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '产品'
        table.cell(0, 1).text = '描述'
        table.cell(1, 0).text = 'A'
        table.cell(1, 1).text = '一|二\n第三行'
        document.add_paragraph('文末说明')
        result = convert(document_bytes(document), '资料.docx')['markdown']
        self.assertLess(result.index('# 资料标题'), result.index('| 产品'))
        self.assertLess(result.index('| 产品'), result.index('文末说明'))
        self.assertIn('一\\|二<br>第三行', result)

    def test_pptx_title_group_and_speaker_notes(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = '唯一标题'
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = '组合内容'
        slide.notes_slide.notes_text_frame.text = '需要保留的演讲备注'
        result = convert(document_bytes(presentation), '会议.pptx')['markdown']
        self.assertEqual(result.count('唯一标题'), 1)
        self.assertIn('组合内容', result)
        self.assertIn('需要保留的演讲备注', result)

    def test_excel_formulas_without_cached_values(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '预算'
        sheet.append(['项目', '金额'])
        sheet.append(['合计', '=SUM(B3:B4)'])
        result = convert(document_bytes(workbook), '预算.xlsx')
        self.assertIn('=SUM(B3:B4)', result['markdown'])
        self.assertIn('公式', result['warning'])

    def test_blank_pdf_is_explicit_error(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaisesRegex(ValueError, 'OCR'):
            convert(data.getvalue(), '扫描.pdf')

    def test_mixed_pdf_flags_pages_without_text(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'):NameObject('/Font'), NameObject('/Subtype'):NameObject('/Type1'), NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 10 100 Td (Readable PDF) Tj ET')
        page[NameObject('/Contents')] = stream
        writer.add_blank_page(width=300, height=300)
        data = io.BytesIO(); writer.write(data)
        result = convert(data.getvalue(), '混合.pdf')
        self.assertIn('Readable PDF', result['markdown'])
        self.assertIn('第 2 页', result['markdown'])
        self.assertIn('部分页面', result['warning'])

    def test_html_removes_scripts_and_keeps_structure(self):
        html = '<head><title>ignore me</title></head><h1>标题</h1><script>bad()</script><style>h1{color:red}</style><p>正文</p>'
        result = convert(html.encode(), '文章.html')['markdown']
        self.assertIn('# 标题', result)
        self.assertIn('正文', result)
        self.assertNotIn('bad()', result)
        self.assertNotIn('ignore me', result)

    def test_tsv_and_csv(self):
        self.assertIn('| 甲 | 乙 |', convert('甲\t乙\n1\t2'.encode(), 'table.tsv')['markdown'])
        self.assertIn('| 甲 | 乙 |', convert('甲,乙\n1,2'.encode(), 'table.csv')['markdown'])

    def test_windowless_logger(self):
        handler = object.__new__(app.AppHandler)
        with patch.object(sys, 'stdout', None):
            handler.log_message('test %s', 'request')


class HttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = app.ThreadingHTTPServer(('127.0.0.1', 0), app.AppHandler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.base = 'http://127.0.0.1:' + str(cls.server.server_port)

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown(); cls.server.server_close(); cls.thread.join(3)

    def test_chinese_multipart_batch_isolates_failure(self):
        body = b''
        for name, data in [('中文.txt', '正常内容'.encode()), ('损坏.pdf', b'bad PDF')]:
            body += ('--testboundary\r\nContent-Disposition: form-data; name="files"; filename="'+name+'"\r\nContent-Type: application/octet-stream\r\n\r\n').encode()+data+b'\r\n'
        body += b'--testboundary--\r\n'
        request = urllib.request.Request(self.base+'/api/convert', data=body, headers={'Content-Type':'multipart/form-data; boundary=testboundary'})
        with urllib.request.urlopen(request) as response:
            result = json.load(response)['results']
        self.assertEqual([r['status'] for r in result], ['success','error'])
        self.assertEqual(result[0]['name'], '中文.txt')

    def test_cross_origin_upload_rejected(self):
        request=urllib.request.Request(self.base+'/api/convert',data=b'x',headers={'Origin':'https://example.com'})
        with self.assertRaises(urllib.error.HTTPError) as caught:
            urllib.request.urlopen(request)
        self.assertEqual(caught.exception.code, 403)
        caught.exception.close()

    def test_static_page_version_and_headers(self):
        with urllib.request.urlopen(self.base) as response:
            self.assertIn('墨转', response.read().decode())
            self.assertIn("script-src 'self'", response.headers['Content-Security-Policy'])
        with urllib.request.urlopen(self.base+'/api/health') as response:
            self.assertEqual(json.load(response)['version'], '2.0.1')


if __name__ == '__main__':
    unittest.main(verbosity=2)

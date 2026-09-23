from __future__ import annotations

import argparse
import base64
import csv
import io
import json
import mimetypes
import re
import sys
import webbrowser
import zipfile
from email import policy
from email.parser import BytesParser
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

APP_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
STATIC_DIR = APP_DIR / "static"
MAX_REQUEST_BYTES = 100 * 1024 * 1024
MAX_FILE_BYTES = 60 * 1024 * 1024
VERSION = "2.0.1"
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".xls", ".csv", ".tsv", ".html", ".htm", ".epub", ".ipynb", ".zip"}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
LEGACY_OFFICE_EXTENSIONS = {".doc", ".ppt"}
PLAIN_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".log",
    ".ini",
    ".cfg",
    ".conf",
    ".env",
    ".yaml",
    ".yml",
    ".toml",
    ".xml",
    ".json",
    ".jsonl",
    ".tex",
}

CODE_LANGUAGES = {
    ".py": "python",
    ".js": "javascript",
    ".mjs": "javascript",
    ".cjs": "javascript",
    ".ts": "typescript",
    ".tsx": "tsx",
    ".jsx": "jsx",
    ".java": "java",
    ".c": "c",
    ".h": "c",
    ".cpp": "cpp",
    ".cc": "cpp",
    ".hpp": "cpp",
    ".cs": "csharp",
    ".go": "go",
    ".rs": "rust",
    ".rb": "ruby",
    ".php": "php",
    ".swift": "swift",
    ".kt": "kotlin",
    ".kts": "kotlin",
    ".scala": "scala",
    ".sh": "bash",
    ".bash": "bash",
    ".zsh": "zsh",
    ".ps1": "powershell",
    ".bat": "batch",
    ".cmd": "batch",
    ".sql": "sql",
    ".css": "css",
    ".scss": "scss",
    ".less": "less",
    ".vue": "vue",
    ".svelte": "svelte",
    ".r": "r",
    ".lua": "lua",
    ".dart": "dart",
}


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")

    for encoding in ("utf-8", "gb18030", "big5", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def tidy_markdown(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return text.strip() + "\n"


def fenced_code(text: str, language: str = "") -> str:
    # A fence longer than any source backtick run preserves nested Markdown/code.
    longest = max((len(match[0]) for match in re.finditer(r"`+", text)), default=0)
    fence = "`" * max(3, longest + 1)
    return f"{fence}{language}\n{text}\n{fence}"


def check_archive(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(item.file_size for item in entries) > 200 * 1024 * 1024:
            raise ValueError("压缩内容过大，请拆分为较小的文件后再转换。")


def markdown_table(rows: list[list[Any]]) -> str:
    cleaned: list[list[str]] = []
    width = 0
    for row in rows:
        values = [
            str(value if value is not None else "")
            .replace("\r", " ")
            .replace("\n", "<br>")
            .replace("|", "\\|")
            for value in row
        ]
        if any(value for value in values):
            cleaned.append(values)
            width = max(width, len(values))
    if not cleaned or width == 0:
        return ""
    cleaned = [row + [""] * (width - len(row)) for row in cleaned]
    header = cleaned[0]
    body = cleaned[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def clean_filename(filename: str) -> str:
    name = Path(filename.replace("\\", "/")).name
    name = re.sub(r"[\x00-\x1f<>:\"/\\|?*]", "_", name).strip(" .")
    return name or "未命名文件"


def title_for(filename: str) -> str:
    return Path(filename).stem.replace("_", " ").strip() or "转换结果"


def image_to_markdown(data: bytes, filename: str, embed_images: bool) -> tuple[str, str]:
    try:
        from PIL import Image
        from io import BytesIO

        image = Image.open(BytesIO(data))
        width, height = image.size
        mode = image.mode
        frame_count = getattr(image, "n_frames", 1)
        details = [
            f"- 尺寸：{width} × {height}",
            f"- 色彩模式：{mode}",
            f"- 帧数：{frame_count}",
        ]
        image.verify()
    except Exception as exc:
        raise ValueError("图片已损坏或无法读取，请重新导出图片后重试。") from exc

    mime = mimetypes.guess_type(filename)[0] or "image/png"
    if embed_images:
        encoded = base64.b64encode(data).decode("ascii")
        image_ref = f"data:{mime};base64,{encoded}"
        warning = "图片已嵌入 Markdown；文件会明显变大，部分 AI 不会读取 data URL 中的图像。"
    else:
        image_ref = filename
        warning = "当前版本没有本地 OCR。请把原图与 Markdown 一起发给支持看图的 AI。"

    markdown = f"""# {title_for(filename)}

![{title_for(filename)}]({image_ref})

## 图片信息

{chr(10).join(details)}

> {warning}
"""
    return tidy_markdown(markdown), warning


def structured_text_to_markdown(data: bytes, filename: str, extension: str) -> str:
    text = decode_text(data).strip()

    if extension == ".json":
        try:
            parsed = json.loads(text)
            text = json.dumps(parsed, ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            pass

    if extension in {".md", ".markdown"}:
        return tidy_markdown(text)
    if extension == ".txt":
        return tidy_markdown(f"# {title_for(filename)}\n\n{text}")

    language = {
        ".json": "json",
        ".jsonl": "jsonl",
        ".yaml": "yaml",
        ".yml": "yaml",
        ".toml": "toml",
        ".xml": "xml",
        ".tex": "latex",
        ".ini": "ini",
        ".cfg": "ini",
        ".conf": "text",
        ".env": "dotenv",
        ".log": "text",
        ".rst": "rst",
        ".txt": "text",
    }.get(extension, "text")

    return tidy_markdown(f"# {title_for(filename)}\n\n{fenced_code(text, language)}")


def code_to_markdown(data: bytes, filename: str, extension: str) -> str:
    text = decode_text(data).rstrip()
    language = CODE_LANGUAGES.get(extension, "")
    return tidy_markdown(f"# {filename}\n\n{fenced_code(text, language)}")


def convert_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF 有密码保护，请先解密后再转换。")
    sections = []
    readable = 0
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            readable += 1
            sections.append(f"## 第 {index} 页\n\n{text}")
        else:
            sections.append(f"## 第 {index} 页\n\n> 未提取到文字：可能是空白页、扫描页或图片页，需对照原文件检查。")
    if not readable:
        raise ValueError("PDF 没有可提取的文字层。扫描版 PDF 需要 OCR（图片文字识别），当前版本尚不支持。")
    return tidy_markdown("\n\n".join(sections))


def convert_docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    chunks: list[str] = []
    content = document.iter_inner_content() if hasattr(document, "iter_inner_content") else document.paragraphs
    for item in content:
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style = (item.style.name if item.style else "").lower()
            match = re.search(r"heading\s*(\d+)", style)
            if match:
                level = min(6, max(1, int(match.group(1))))
                chunks.append(f"{'#' * level} {text}")
            elif style.startswith("title"):
                chunks.append(f"# {text}")
            elif style.startswith("list"):
                chunks.append(f"- {text}")
            else:
                chunks.append(text)
        elif isinstance(item, Table):
            table = markdown_table([[cell.text.strip() for cell in row.cells] for row in item.rows])
            if table:
                chunks.append(table)
    if not hasattr(document, "iter_inner_content"):
        for table_item in document.tables:
            table = markdown_table([[cell.text.strip() for cell in row.cells] for row in table_item.rows])
            if table:
                chunks.append(table)
    return tidy_markdown("\n\n".join(chunks))


def convert_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    chunks: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = ""
        if slide.shapes.title and getattr(slide.shapes.title, "text", ""):
            title = slide.shapes.title.text.strip()
        chunks.append(f"## 幻灯片 {slide_index}" + (f"：{title}" if title else ""))
        def walk_shapes(shapes):
            for child in shapes:
                if hasattr(child, "shapes"):
                    yield from walk_shapes(child.shapes)
                else:
                    yield child
        for shape in walk_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if getattr(shape, "has_table", False):
                table = markdown_table([[cell.text.strip() for cell in row.cells] for row in shape.table.rows])
                if table:
                    chunks.append(table)
            elif getattr(shape, "has_text_frame", False):
                lines = [paragraph.text.strip() for paragraph in shape.text_frame.paragraphs if paragraph.text.strip()]
                if lines:
                    chunks.append("\n\n".join(lines))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                chunks.append("### 演讲备注\n\n" + notes.text.strip())
    return tidy_markdown("\n\n".join(chunks))


def convert_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    # Formula text is retained: cached values are often absent in exported workbooks.
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    chunks: list[str] = []
    for sheet in workbook.worksheets:
        if (sheet.max_row or 0) * (sheet.max_column or 0) > 1000000:
            workbook.close()
            raise ValueError("工作表范围超过 100 万个单元格，请拆分表格或清除空白区域的格式后重试。")
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        table = markdown_table(rows)
        if table:
            chunks.append(f"## {sheet.title}\n\n{table}")
    workbook.close()
    return tidy_markdown("\n\n".join(chunks))


def convert_xls(data: bytes) -> str:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=data)
    chunks: list[str] = []
    for sheet in workbook.sheets():
        rows = [sheet.row_values(index) for index in range(sheet.nrows)]
        table = markdown_table(rows)
        if table:
            chunks.append(f"## {sheet.name}\n\n{table}")
    return tidy_markdown("\n\n".join(chunks))


def convert_csv(data: bytes) -> str:
    text = decode_text(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return tidy_markdown(markdown_table(rows))


def convert_html(data: bytes) -> str:
    from markdownify import markdownify
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(decode_text(data), "html.parser")
    for element in soup(["script", "style", "head", "noscript", "iframe", "template"]):
        element.decompose()
    return tidy_markdown(markdownify(str(soup), heading_style="ATX"))


def convert_ipynb(data: bytes, filename: str) -> str:
    notebook = json.loads(decode_text(data))
    chunks = [f"# {title_for(filename)}"]
    for cell in notebook.get("cells", []):
        source = "".join(cell.get("source", [])).strip()
        if not source:
            continue
        if cell.get("cell_type") == "markdown":
            chunks.append(source)
        elif cell.get("cell_type") == "code":
            language = notebook.get("metadata", {}).get("kernelspec", {}).get("language", "")
            chunks.append(fenced_code(source, language))
            outputs = []
            for output in cell.get("outputs", []):
                output_text = output.get("text") or output.get("data", {}).get("text/plain")
                if output_text:
                    outputs.append("".join(output_text) if isinstance(output_text, list) else str(output_text))
            if outputs:
                chunks.append(fenced_code("\n".join(outputs).strip(), "text"))
    return tidy_markdown("\n\n".join(chunks))


def convert_epub(data: bytes) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = [name for name in archive.namelist() if Path(name).suffix.lower() in {".html", ".htm", ".xhtml"}]
        for name in names:
            try:
                content = convert_html(archive.read(name)).strip()
            except Exception:
                continue
            if content:
                chunks.append(f"## {Path(name).stem}\n\n{content}")
    return tidy_markdown("\n\n".join(chunks))


def convert_zip(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = [entry for entry in archive.infolist() if not entry.is_dir()]
        rows = [["文件", "大小（字节）", "压缩后大小"]]
        rows.extend([[entry.filename, entry.file_size, entry.compress_size] for entry in entries])
        chunks = ["# 压缩包内容", markdown_table(rows)]
        for entry in entries:
            extension = Path(entry.filename).suffix.lower()
            if extension in PLAIN_TEXT_EXTENSIONS or extension in CODE_LANGUAGES or extension in {".csv", ".html", ".htm"}:
                if entry.file_size > 2 * 1024 * 1024:
                    continue
                content = archive.read(entry)
                if extension == ".csv":
                    converted = convert_csv(content)
                elif extension in {".html", ".htm"}:
                    converted = convert_html(content)
                elif extension in CODE_LANGUAGES:
                    converted = code_to_markdown(content, entry.filename, extension)
                else:
                    converted = structured_text_to_markdown(content, entry.filename, extension)
                chunks.append(f"## {entry.filename}\n\n{converted.strip()}")
    return tidy_markdown("\n\n".join(chunks))


def convert_document(data: bytes, filename: str, extension: str) -> str:
    converters = {
        ".pdf": lambda: convert_pdf(data),
        ".docx": lambda: convert_docx(data),
        ".pptx": lambda: convert_pptx(data),
        ".xlsx": lambda: convert_xlsx(data),
        ".xls": lambda: convert_xls(data),
        ".csv": lambda: convert_csv(data),
        ".tsv": lambda: tidy_markdown(markdown_table(list(csv.reader(io.StringIO(decode_text(data)), delimiter="\t")))),
        ".html": lambda: convert_html(data),
        ".htm": lambda: convert_html(data),
        ".ipynb": lambda: convert_ipynb(data, filename),
        ".epub": lambda: convert_epub(data),
        ".zip": lambda: convert_zip(data),
    }
    converter = converters.get(extension)
    if converter is None:
        raise ValueError(f"暂不支持这种格式：{extension or '无扩展名'}")
    return converter()


def convert_file(
    data: bytes,
    filename: str,
    *,
    include_metadata: bool,
    embed_images: bool,
) -> dict[str, Any]:
    safe_name = clean_filename(filename)
    extension = Path(safe_name).suffix.lower()
    warning = ""

    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"单个文件不能超过 {MAX_FILE_BYTES // (1024 * 1024)} MB")
    if not data:
        raise ValueError("文件是空的")
    if extension in LEGACY_OFFICE_EXTENSIONS:
        raise ValueError("旧版 .doc/.ppt 暂不支持，请先用 Office 另存为 .docx/.pptx")
    if extension in {".docx", ".pptx", ".xlsx", ".epub", ".zip"}:
        try:
            check_archive(data)
        except zipfile.BadZipFile as exc:
            raise ValueError("文件已损坏或扩展名与实际格式不符，请从原软件重新导出。") from exc

    if extension in IMAGE_EXTENSIONS:
        markdown, warning = image_to_markdown(data, safe_name, embed_images)
    elif extension in CODE_LANGUAGES:
        markdown = code_to_markdown(data, safe_name, extension)
    elif extension in PLAIN_TEXT_EXTENSIONS:
        markdown = structured_text_to_markdown(data, safe_name, extension)
    elif extension in DOCUMENT_EXTENSIONS:
        try:
            markdown = convert_document(data, safe_name, extension)
        except ValueError:
            raise
        except Exception as exc:
            if isinstance(exc, ImportError):
                raise ValueError(f"解析 {extension} 的组件缺失，请重新解压完整的便携包。") from exc
            raise ValueError(f"无法解析 {extension} 文件，可能已损坏、加密或格式不符，请重新导出后重试。") from exc
    else:
        # Unknown binary formats must not be turned into plausible-looking gibberish.
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError(f"暂不支持 {extension or '无扩展名'}，请先导出为常见文档或 UTF-8 文本。") from exc
        if any(not char.isprintable() and char not in "\n\r\t" for char in text):
            raise ValueError(f"暂不支持这种二进制格式：{extension or '无扩展名'}")
        markdown = structured_text_to_markdown(data, safe_name, extension)
        warning = f"扩展名 {extension or '（无）'} 未识别，已按 UTF-8 文本读取。"

    if extension == ".pdf":
        warning = "PDF 按页提取文字，复杂表格和多栏排版请对照原文件检查。"
        if "未提取到文字：" in markdown:
            warning += " 部分页面未提取到文字，已在结果中标记。"
    elif extension == ".xlsx":
        warning = "表格中的公式保留为公式文本；图表与内嵌图片未提取。"
    elif extension in {".docx", ".pptx"}:
        warning = "已提取正文和表格；内嵌图片及图表中的信息请对照原文件检查。"
    elif extension == ".zip":
        warning = "已生成文件清单，并提取不超过 2 MB 的文本、代码、CSV 和 HTML；其他文件请解压后单独导入。"

    if not markdown.strip():
        raise ValueError("没有提取到可读内容；如果这是扫描 PDF，请改为上传原图或使用 OCR")

    if include_metadata:
        metadata = (
            f"---\n"
            f"source_file: {json.dumps(safe_name, ensure_ascii=False)}\n"
            f"source_type: {json.dumps(extension.lstrip('.') or 'unknown', ensure_ascii=False)}\n"
            f"source_size: {len(data)}\n"
            f"---\n\n"
        )
        markdown = metadata + markdown

    output_name = f"{Path(safe_name).stem}.md"
    return {
        "name": safe_name,
        "outputName": output_name,
        "markdown": markdown,
        "characters": len(markdown),
        "warning": warning,
        "status": "success",
    }


def parse_multipart(content_type: str, body: bytes) -> tuple[list[tuple[str, bytes]], dict[str, str]]:
    message = BytesParser(policy=policy.default).parsebytes(
        b"Content-Type: "
        + content_type.encode("utf-8")
        + b"\r\nMIME-Version: 1.0\r\n\r\n"
        + body
    )
    files: list[tuple[str, bytes]] = []
    fields: dict[str, str] = {}

    if not message.is_multipart():
        return files, fields

    for part in message.iter_parts():
        if part.get_content_disposition() != "form-data":
            continue
        field_name = part.get_param("name", header="content-disposition") or ""
        filename = part.get_filename()
        payload = part.get_payload(decode=True) or b""
        if filename:
            files.append((filename, payload))
        else:
            charset = part.get_content_charset() or "utf-8"
            fields[field_name] = payload.decode(charset, errors="replace")
    return files, fields


class AppHandler(BaseHTTPRequestHandler):
    server_version = "FormatToMarkdown/2.0"

    def end_headers(self) -> None:
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header("Content-Security-Policy", "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; object-src 'none'; base-uri 'none'; frame-ancestors 'none'")
        super().end_headers()

    def log_message(self, format: str, *args: Any) -> None:
        if sys.stdout is not None:
            print(f"[{self.log_date_time_string()}] {format % args}")

    def send_json(self, payload: dict[str, Any], status: int = HTTPStatus.OK) -> None:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        request_path = unquote(parsed.path)
        if request_path == "/favicon.ico":
            self.send_response(HTTPStatus.NO_CONTENT)
            self.end_headers()
            return
        if request_path == "/api/health":
            self.send_json({"ok": True, "service": "format-to-markdown", "version": VERSION})
            return

        relative = "index.html" if request_path in {"/", ""} else request_path.lstrip("/")
        candidate = (STATIC_DIR / relative).resolve()
        try:
            candidate.relative_to(STATIC_DIR.resolve())
        except ValueError:
            self.send_error(HTTPStatus.FORBIDDEN)
            return

        if not candidate.is_file():
            self.send_error(HTTPStatus.NOT_FOUND)
            return

        content = candidate.read_bytes()
        content_type = mimetypes.guess_type(candidate.name)[0] or "application/octet-stream"
        if content_type.startswith("text/") or candidate.suffix in {".js", ".json"}:
            content_type += "; charset=utf-8"
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(content)))
        self.send_header("Cache-Control", "no-cache")
        self.end_headers()
        self.wfile.write(content)

    def do_POST(self) -> None:
        origin = self.headers.get("Origin")
        if origin and origin != f"http://{self.headers.get('Host')}":
            self.send_json({"error": "请从本工具的窗口内操作。"}, HTTPStatus.FORBIDDEN)
            return
        if urlparse(self.path).path != "/api/convert":
            self.send_error(HTTPStatus.NOT_FOUND)
            return

        try:
            content_length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            self.send_json({"error": "无效的请求大小"}, HTTPStatus.BAD_REQUEST)
            return

        if content_length <= 0 or content_length > MAX_REQUEST_BYTES:
            self.send_json(
                {"error": f"单次上传总大小不能超过 {MAX_REQUEST_BYTES // (1024 * 1024)} MB"},
                HTTPStatus.REQUEST_ENTITY_TOO_LARGE,
            )
            return

        content_type = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in content_type:
            self.send_json({"error": "请使用文件上传格式"}, HTTPStatus.BAD_REQUEST)
            return

        try:
            body = self.rfile.read(content_length)
            files, fields = parse_multipart(content_type, body)
            if len(files) > 50:
                self.send_json({"error": "每批最多 50 个文件，请分批转换。"}, HTTPStatus.BAD_REQUEST)
                return
            if not files:
                self.send_json({"error": "没有收到文件"}, HTTPStatus.BAD_REQUEST)
                return

            include_metadata = fields.get("includeMetadata", "false").lower() == "true"
            embed_images = fields.get("embedImages", "false").lower() == "true"
            results = []
            for filename, data in files:
                try:
                    results.append(
                        convert_file(
                            data,
                            filename,
                            include_metadata=include_metadata,
                            embed_images=embed_images,
                        )
                    )
                except Exception as exc:
                    safe_name = clean_filename(filename)
                    results.append(
                        {
                            "name": safe_name,
                            "outputName": f"{Path(safe_name).stem}.md",
                            "markdown": "",
                            "characters": 0,
                            "warning": "",
                            "status": "error",
                            "error": str(exc),
                        }
                    )

            self.send_json({"results": results})
        except Exception as exc:
            self.send_json({"error": f"转换请求处理失败：{exc}"}, HTTPStatus.INTERNAL_SERVER_ERROR)


def main() -> None:
    parser = argparse.ArgumentParser(description="把常见文件转换为 Markdown")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--no-browser", action="store_true")
    args = parser.parse_args()

    if not STATIC_DIR.is_dir():
        raise SystemExit(f"找不到界面文件：{STATIC_DIR}")

    server = ThreadingHTTPServer((args.host, args.port), AppHandler)
    url = f"http://{args.host}:{args.port}"
    print(f"格式转 Markdown 已启动：{url}")
    print("按 Ctrl+C 停止。")
    if not args.no_browser:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n已停止。")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
